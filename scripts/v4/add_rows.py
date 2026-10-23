"""Add drafted rows to data/Routing-Table_V4.xlsx, keeping its formatting and dropdowns.

    python3 scripts/v4/add_rows.py rows/*.json

Each file is { "items": [...], "sources": [...], "safety": [...], "overlap_additions": [...] } with
columns named as the workbook's headers (snake_case). Items and sources whose key or id already
exists are updated in place; new ones are appended. An overlap addition joins its members to an
existing group, or creates the group when it carries a rule. Then run `pnpm catalog:build`.
"""

import copy
import json
import pathlib
import sys

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
BOOK = ROOT / "data" / "Routing-Table_V4.xlsx"


def as_cell(value):
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if value is None or value == "":
        return None
    return value


def upsert(ws, key_column, rows):
    header = [c.value for c in ws[1]]
    index = {ws.cell(row=r, column=1 + header.index(key_column)).value: r for r in range(2, ws.max_row + 1)}
    style_row = 2 if ws.max_row >= 2 else None
    added = updated = 0
    for row in rows:
        unknown = set(row) - set(header)
        if unknown:
            raise SystemExit(f"{ws.title}: unknown column(s) {sorted(unknown)} in {row.get(key_column)}")
        key = row[key_column]
        if key in index:
            r = index[key]
            updated += 1
        else:
            r = ws.max_row + 1
            index[key] = r
            added += 1
        for c, name in enumerate(header, start=1):
            cell = ws.cell(row=r, column=c)
            if style_row and r != style_row:
                src = ws.cell(row=style_row, column=c)
                cell.font, cell.alignment, cell.border, cell.fill = (
                    copy.copy(src.font), copy.copy(src.alignment), copy.copy(src.border), copy.copy(src.fill))
                cell.number_format = src.number_format
            if name in row:
                cell.value = as_cell(row[name])
    return added, updated


def main(paths):
    wb = load_workbook(BOOK)
    items, sources, safety, overlaps = [], [], [], []
    for p in paths:
        data = json.loads(pathlib.Path(p).read_text())
        items += data.get("items", [])
        sources += data.get("sources", [])
        safety += data.get("safety", [])
        overlaps += data.get("overlap_additions", [])

    print("Items", upsert(wb["Items"], "key", items))
    # The same paper drafted by two batches keeps the first copy.
    seen, unique = set(), []
    for s in sources:
        # A bare DOI ("10.2196/12053") becomes a link the validator accepts.
        url = str(s.get("doi_or_url") or "").strip()
        if url.lower().startswith("doi:"):
            url = url[4:].strip()
        if url.startswith("10."):
            s["doi_or_url"] = "https://doi.org/" + url
        if s["id"] not in seen:
            seen.add(s["id"])
            unique.append(s)
    print("Sources", upsert(wb["Sources"], "id", unique))

    ws = wb["Safety"]
    header = [c.value for c in ws[1]]
    existing = {(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value) for r in range(2, ws.max_row + 1)}
    fresh = [s for s in safety if (s["item"], s["condition"]) not in existing]
    for s in fresh:
        ws.append([as_cell(s.get(h)) for h in header])
    print("Safety added", len(fresh))

    ws = wb["Overlaps"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    rows = {ws.cell(row=r, column=col["group"]).value: r for r in range(2, ws.max_row + 1)}
    for o in overlaps:
        group = o["group"]
        if group in rows:
            cell = ws.cell(row=rows[group], column=col["members"])
            members = [m.strip() for m in (cell.value or "").split(",") if m.strip()]
            for m in o["members"]:
                if m not in members:
                    members.append(m)
            cell.value = ", ".join(members)
        elif o.get("rule"):
            ws.append([as_cell({"group": group, "members": o["members"], "rule": o["rule"], "note": o.get("note")}.get(h)) for h in header])
            rows[group] = ws.max_row
        else:
            print(f"skipped overlap addition for unknown group {group} (no rule given)")
    wb.save(BOOK)
    print("saved", BOOK)


if __name__ == "__main__":
    main(sys.argv[1:])
