"""Write data/Routing-Table_V4.xlsx from a JSON draft (one-time bootstrap).

After this, the workbook is the source of truth: edit it in Google Sheets (File > Import, then
File > Download > xlsx back into data/), and `pnpm catalog:build` rebuilds data/catalog-v4.json.

    python3 scripts/v4/seed_workbook.py scripts/v4/draft.json

The draft is { "items": [...], "sources": [...], "overlaps": [...], "safety": [...], "banned": [...] }
with item fields named as in packages/catalog/src/ingest-v4.ts (snake_case headers).
"""

import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "Routing-Table_V4.xlsx"

ITEM_COLUMNS = [
    "key", "name", "short_name", "area", "also_serves", "kind", "moments", "cadence", "ritual_role",
    "cost_month_usd", "how_common", "comes_from", "why_loved", "evidence_tier", "evidence_for",
    "evidence_requires", "tier_otherwise", "evidence_summary", "sources", "dose_min", "dose_typical",
    "dose_unit", "studied_forms", "keep_when", "clear_when", "never_clear", "human_question",
    "fits_instead", "awe", "wearable_sees", "overlaps", "top50", "status", "reviewed_by",
    "reviewed_on", "notes_internal",
]
SOURCE_COLUMNS = ["id", "citation", "year", "type", "doi_or_url", "finding", "population", "checked_by", "checked_on"]
OVERLAP_COLUMNS = ["group", "members", "rule", "note"]
SAFETY_COLUMNS = ["item", "condition", "action", "note"]

# One value from a list: a dropdown. Several values: a note listing what is allowed.
SINGLE = {
    "area": ["eat", "move", "mind", "sleep", "skin"],
    "kind": ["supplement", "food-drink", "movement", "practice", "device", "service", "environment", "timing"],
    "cadence": ["daily", "weekly", "monthly", "occasional"],
    "ritual_role": ["anchor", "companion", "standalone"],
    "evidence_tier": ["A", "B", "C", "D", "N"],
    "evidence_requires": ["none", "trains", "deficiency", "low_intake"],
    "tier_otherwise": ["A", "B", "C", "D", "N"],
    "dose_unit": ["g", "mg", "mcg", "IU", "min", "sessions/week"],
    "never_clear": ["yes", "no"],
    "awe": ["nature", "music", "visual", "collective", "moral", "spiritual", "life", "epiphany"],
    "wearable_sees": ["yes", "partial", "no"],
    "top50": ["yes", "no"],
    "status": ["draft", "reviewed", "live"],
    "type": ["meta-analysis", "RCT", "position stand", "review", "cohort", "expert"],
    "rule": ["keep-one", "keep-best-evidence", "keep-most-used", "show-together"],
    "condition": [
        "any", "on-medication", "pregnant", "breastfeeding", "kidney-disease", "liver-disease",
        "heart-condition", "blood-pressure", "diabetes", "under-18", "skin-condition", "bleeding-risk",
        "mental-health-condition",
    ],
    "action": ["clinician-first", "see-dermatologist", "never-suggest", "note"],
}
MULTI = {
    "also_serves": "Up to two of: eat, move, mind, sleep, skin. Separate with commas.",
    "moments": "One or more of: morning, midday, evening, night, weekend, any.",
    "comes_from": "One or more of: feed, ad, podcast, friend, gym, clinician, brand, culture.",
    "why_loved": "Up to three of: connection, ritual, confidence, looks, calm, energy, belonging, treat.",
    "keep_when": "Codes: THEIRS, LOVED (always), EVIDENCE_REGULAR (tier A or B), TRAINS.",
    "clear_when": "Codes: BORROWED_WEAK, OVERLAP, UNUSED, DOSE_BELOW_MIN, FORM_NOT_ABSORBED, UNLOVED_HABIT, NOT_TRAINING, TOO_MANY_STEPS.",
    "sources": "Source ids from the Sources sheet, separated by commas.",
    "studied_forms": "Forms used in the studies, separated by commas.",
    "overlaps": "Group names from the Overlaps sheet, separated by commas.",
    "members": "Item keys, separated by commas.",
}
WIDE = {"evidence_summary": 60, "evidence_for": 40, "human_question": 40, "fits_instead": 40,
        "notes_internal": 50, "citation": 60, "finding": 50, "note": 50, "name": 32, "key": 34,
        "doi_or_url": 40, "population": 30}

HEAD_FILL = PatternFill("solid", fgColor="11191B")
HEAD_FONT = Font(bold=True, color="F0ECE2")


def cell(value):
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    if value is None:
        return None
    return value


def sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for c in ws[1]:
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([cell(r.get(c)) for c in columns])
    ws.freeze_panes = "B2"
    last = max(ws.max_row, 2) + 300
    for i, name in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = WIDE.get(name, 16)
        if name in SINGLE:
            dv = DataValidation(type="list", formula1='"' + ",".join(SINGLE[name]) + '"', allow_blank=True)
            dv.error = "Pick one of the listed values."
            dv.errorTitle = name
            dv.prompt = "Pick one: " + ", ".join(SINGLE[name])
            dv.promptTitle = name
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{last}")
        elif name in MULTI:
            dv = DataValidation(type="custom", formula1="TRUE", allow_blank=True)
            dv.prompt = MULTI[name]
            dv.promptTitle = name
            dv.showErrorMessage = False
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{last}")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


README = [
    ("Routing table V4", ""),
    ("What this is", "One row per thing a person does or buys for their health. The engine reads it to decide what to protect, what to propose letting go of, and where each thing sits in a person's Ritual."),
    ("Who edits it", "Ishan owns every call in it. Claude drafted rows from sources it opened; nothing reaches people until a row's status is live, which needs reviewed_by and reviewed_on."),
    ("How to edit", "Import into Google Sheets, edit, then File > Download > Microsoft Excel and save over data/Routing-Table_V4.xlsx. Run pnpm catalog:build; pnpm check fails with the sheet, row and column of anything wrong."),
    ("Columns move freely", "The build reads columns by their header text, so reorder them as you like. Never rename a header."),
    ("Tier A", "Meta-analyses or several consistent RCTs in healthy adults, for evidence_for. Protects it when done regularly; never called noise, even when unfelt."),
    ("Tier B", "RCTs agree but are small, short or in a narrower group. Protects it when done regularly."),
    ("Tier C", "Studies disagree, or are small and short. Neither protects nor clears; their values decide."),
    ("Tier D", "Mechanism, animal studies or anecdote only. Can be cleared when borrowed and unloved."),
    ("Tier N", "No independent studies (brand-funded only counts as none). Can be cleared when borrowed and unloved."),
    ("evidence_requires", "When the tier only holds for some people (creatine needs training), name the condition and give tier_otherwise for everyone else."),
    ("Direction", "keep_when with EVIDENCE_REGULAR means the evidence is of a benefit and protects. A clear_when of UNLOVED_HABIT means the evidence is of a cost (a nightcap, nicotine), and it protects nothing."),
    ("Sources", "A and B rows need at least one source with checked_by and checked_on filled. Only list a source someone actually opened. Findings are paraphrased, 30 words or fewer."),
    ("Words a person reads", "evidence_summary (40 words or fewer), human_question (under 12 words) and fits_instead never use the words on the Words sheet, never say they or their, and never end in an exclamation mark."),
]


def main(path):
    draft = json.loads(pathlib.Path(path).read_text())
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    for k, v in README:
        readme.append([k, v])
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 110
    for row in readme.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    sheet(wb, "Items", ITEM_COLUMNS, draft["items"])
    sheet(wb, "Sources", SOURCE_COLUMNS, draft["sources"])
    sheet(wb, "Overlaps", OVERLAP_COLUMNS, draft["overlaps"])
    sheet(wb, "Safety", SAFETY_COLUMNS, draft["safety"])
    sheet(wb, "Words", ["phrase"], [{"phrase": p} for p in draft["banned"]])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT.relative_to(ROOT)}: {len(draft['items'])} items, {len(draft['sources'])} sources, "
          f"{len(draft['overlaps'])} overlap groups, {len(draft['safety'])} safety rows")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else str(ROOT / "scripts" / "v4" / "draft.json"))
