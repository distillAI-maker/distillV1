"""Step 0: pull the Items and Sources sheets out of the workbook into DIR/catalog.json for the checks.

    python3 scripts/v4/validate/extract.py DIR      (needs openpyxl)
"""
import json, re, sys, pathlib
from openpyxl import load_workbook

DIR = sys.argv[1]
BOOK = pathlib.Path(__file__).resolve().parents[3] / 'data' / 'Routing-Table_V4.xlsx'
wb = load_workbook(BOOK, read_only=True)

def rows(name):
    h, out = None, []
    for r in wb[name].iter_rows(values_only=True):
        cells = [str(c).strip() if c else '' for c in r]
        if h is None:
            if 'key' in cells or 'id' in cells: h = cells
            continue
        if any(r): out.append({k: (v if v is not None else '') for k, v in zip(h, r)})
    return out

items, src = rows('Items'), rows('Sources')
used = {}
for i in items:
    for x in re.split(r'[,; ]+', str(i['sources'])):
        if x: used.setdefault(x, []).append(i['key'])
for s in src: s['rows'] = used.get(s['id'], [])
pathlib.Path(DIR).mkdir(parents=True, exist_ok=True)
json.dump({'items': items, 'sources': src}, open(f'{DIR}/catalog.json', 'w'), default=str, indent=1)
print(len(items), 'items,', len(src), 'sources')
